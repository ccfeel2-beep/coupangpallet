from flask import Flask, render_template, request, send_file
from openpyxl import load_workbook
from pptx import Presentation
from copy import deepcopy
from pathlib import Path
from io import BytesIO
import re

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024

REQUIRED_HEADERS = ['파렛트', 'SKU No', 'SKU NAME', 'box수량', '수량', '제조일자']


def norm(v):
    return re.sub(r'\s+', '', str(v or '')).lower()


def cell_text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_excel(file_obj):
    wb = load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb.active
    meta = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True):
        if len(row) >= 2 and row[0] is not None:
            key = norm(row[0])
            if key in ('입고예정일자', '납품센터명', '발주센터명', '발주번호'):
                meta[key] = cell_text(row[1])
    header_row = None
    headers = []
    for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        normalized = [norm(x) for x in row]
        if all(norm(h) in normalized for h in REQUIRED_HEADERS):
            header_row, headers = ridx, list(row)
            break
    if not header_row:
        raise ValueError('엑셀에서 상품 표의 열 제목을 찾지 못했습니다. 필요한 열: ' + ', '.join(REQUIRED_HEADERS))
    idx = {norm(h): i for i, h in enumerate(headers) if h is not None}
    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(v is not None for v in row):
            continue
        pallet = cell_text(row[idx[norm('파렛트')]])
        if not pallet:
            continue
        rec = {h: (row[idx[norm(h)]] if idx[norm(h)] < len(row) else None) for h in REQUIRED_HEADERS}
        rec['파렛트'] = pallet
        records.append(rec)
    if not records:
        raise ValueError('파렛트가 지정된 상품 행이 없습니다.')
    missing = [k for k in ('입고예정일자','납품센터명','발주센터명','발주번호') if not meta.get(k)]
    if missing:
        raise ValueError('기본 정보가 비어 있습니다: ' + ', '.join(missing))
    return meta, records


def replace_text(shape, new_text):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    # Keep first run's formatting where possible; preserve paragraph structure by placing text in first run.
    if not tf.paragraphs:
        tf.text = new_text
        return
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = new_text
        for run in p.runs[1:]:
            run.text = ''
    else:
        p.text = new_text
    for para in tf.paragraphs[1:]:
        for run in para.runs:
            run.text = ''


def set_cell(cell, value):
    text = cell_text(value)
    tf = cell.text_frame
    if not tf.paragraphs:
        cell.text = text
        return
    p = tf.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.text = text
    for para in tf.paragraphs[1:]:
        for r in para.runs:
            r.text = ''


def clone_slide(prs, source_slide):
    slide = prs.slides.add_slide(source_slide.slide_layout)
    # Remove default placeholders on the new slide.
    for shp in list(slide.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in source_slide.shapes:
        newel = deepcopy(shp.element)
        slide.shapes._spTree.insert_element_before(newel, 'p:extLst')
    # Copy relationships for images/other linked elements if any.
    for rel in source_slide.part.rels.values():
        if 'notesSlide' not in rel.reltype and 'slideLayout' not in rel.reltype:
            try:
                slide.part.rels.add_relationship(rel.reltype, rel._target, rel.rId, is_external=rel.is_external)
            except Exception:
                pass
    return slide


def rebuild_table(table, records):
    # Header row remains; use existing data row as a formatting template.
    while len(table.rows) > 2:
        table._tbl.remove(table.rows[len(table.rows) - 1]._tr)
    if not records:
        records = []
    template_tr = deepcopy(table.rows[1]._tr)
    if records:
        # First data row will be reused, subsequent rows clone the row XML to preserve formatting.
        for i, rec in enumerate(records):
            if i == 0:
                row = table.rows[1]
            else:
                # Insert cloned row before any trailing extension elements, preserving table XML order.
                new_tr = table._tbl.add_tr(height=template_tr.h)
                for tc in list(new_tr.tc_lst):
                    new_tr.remove(tc)
                for tc in template_tr.tc_lst:
                    new_tr.append(deepcopy(tc))
                row = table.rows[len(table.rows) - 1]
            vals = [i + 1, rec['SKU No'], rec['SKU NAME'], rec['box수량'], rec['수량'], rec['제조일자']]
            for c, val in zip(row.cells, vals):
                set_cell(c, val)
    else:
        for c, val in zip(table.rows[1].cells, ['', '', '자료 없음', '', '', '']):
            set_cell(c, val)


def make_ppt(template_path, meta, records):
    prs = Presentation(template_path)
    if not prs.slides:
        raise ValueError('PPT 템플릿에 슬라이드가 없습니다.')
    # Use the first existing slide as the design template, and remove the rest.
    first = prs.slides[0]
    slide_id_list = prs.slides._sldIdLst
    for sid in list(slide_id_list)[1:]:
        prs.part.drop_rel(sid.rId)
        slide_id_list.remove(sid)
    grouped = {}
    for rec in records:
        grouped.setdefault(rec['파렛트'], []).append(rec)
    # Sort numeric pallet numbers numerically, then nonnumeric labels alphabetically.
    keys = sorted(grouped, key=lambda x: (0, int(x)) if str(x).isdigit() else (1, str(x)))
    total = len(keys)
    for n, key in enumerate(keys, start=1):
        slide = first if n == 1 else clone_slide(prs, first)
        box_total = 0
        for r in grouped[key]:
            try:
                box_total += float(str(r['box수량']).replace(',', '').strip() or 0)
            except (TypeError, ValueError):
                raise ValueError(f'파렛트 {key}의 BOX 수량이 숫자가 아닙니다: {r["box수량"]}')
        if box_total.is_integer():
            box_label = str(int(box_total))
        else:
            box_label = str(box_total)
        for shape in slide.shapes:
            if shape.has_text_frame:
                if shape.text.startswith('  6 - 1') or ('총 박스수량' in shape.text and 'BOX' in shape.text):
                    replace_text(shape, f'  {total} - {n}                     / 총 박스수량  (   {box_label}   BOX)')
                elif '입고예정일자' in shape.text and '발주번호' in shape.text:
                    text = (f'입고예정일자 (   {meta["입고예정일자"]}  )   / 납품센터명 (      {meta["납품센터명"]}  )\n'
                            f'업체명    ( 주식회사 봉우리   )\t  / 발주센터명 (\t {meta["발주센터명"]} )\n'
                            f'발주번호      ({meta["발주번호"]} )')
                    replace_text(shape, text)
            if shape.has_table:
                rebuild_table(shape.table, grouped[key])
    out = BytesIO()
    prs.save(out)
    out.seek(0)
    return out, total, len(records)


@app.get('/')
def index():
    return render_template('index.html')


@app.post('/generate')
def generate():
    excel = request.files.get('excel')
    template = request.files.get('template')
    if not excel or not excel.filename.lower().endswith(('.xlsx', '.xlsm')):
        return render_template('index.html', error='엑셀 파일(.xlsx 또는 .xlsm)을 선택해 주세요.'), 400
    if not template or not template.filename.lower().endswith('.pptx'):
        return render_template('index.html', error='PPT 양식 파일(.pptx)을 선택해 주세요.'), 400
    try:
        meta, records = read_excel(excel)
        result, pallet_count, row_count = make_ppt(template, meta, records)
        safe_date = re.sub(r'[^0-9가-힣]', '', meta['입고예정일자']) or '입고예정'
        return send_file(result, as_attachment=True, download_name=f'쿠팡_파렛트적재리스트_{safe_date}.pptx', mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')
    except Exception as e:
        return render_template('index.html', error=f'파일을 생성하지 못했습니다: {e}'), 400

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=False)
